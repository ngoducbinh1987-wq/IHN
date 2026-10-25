"""Kiểm thử đối chiếu database với file Excel bản 3.0.

1. Định mức: kg/km đồng, kg/km cách điện, Rdc của 78 SKU do master.calc_spec tính phải khớp Dinh_muc.
2. MRP: tổng cấp SX từng vật tư do plan.run_mrp tính phải khớp công thức BOM_chi_tiet
   (H = D·E/(1−hao hụt) + I, với I là phế khởi động theo công thức gốc của file).

Dùng: python3 -I tests/test_parity.py <file.xlsx> "<chuỗi kết nối psql>"
"""
import json
import math
import time
import subprocess
import sys

import openpyxl

XLSX, PG = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(XLSX, data_only=True)
TOL = 1e-6


def psql(sql):
    r = subprocess.run(["psql", PG, "-v", "ON_ERROR_STOP=1", "-qAt", "-c", sql], capture_output=True, text=True)
    if r.returncode:
        raise SystemExit(r.stderr)
    return r.stdout.strip()


fails = 0

# ------------------------------------------------------------------ 1. định mức
dm = {r[0]: r for r in wb["Dinh_muc"].iter_rows(min_row=9, max_row=86, max_col=43, values_only=True)}
db = json.loads(psql("""select json_object_agg(p.code, json_build_array(r.metal_kgkm, r.ins_kgkm, r.rdc_calc))
  from master.spec_results r join master.product_specs s on s.id = r.spec_id join master.products p on p.id = s.product_id
  where s.version = 1"""))
worst = 0.0
for sku, (metal, ins, rdc) in db.items():
    x = dm[sku]
    for got, exp, name in ((metal, x[22], "kim loại kg/km"), (ins, x[23], "cách điện kg/km"), (rdc, x[29], "Rdc")):
        err = abs(got / exp - 1)
        worst = max(worst, err)
        if err > TOL:
            fails += 1
            print(f"SAI {sku} {name}: DB {got} ≠ Excel {exp}")
print(f"[1] Định mức {len(db)} SKU × 3 chỉ tiêu: sai số tương đối lớn nhất {worst:.2e}")

# ------------------------------------------------------------------ 2. MRP
P = {r[0]: float(r[2]) for r in wb["Tham_so"].iter_rows(min_row=9, max_row=67, max_col=5, values_only=True) if r[0]}
mats = {}
ns = wb["Nguon_spec"]
for r in ns.iter_rows(min_row=117, max_row=167, max_col=5, values_only=True):
    if r[0]:
        mats[r[0]] = float(r[4] or 0)
bom = [r for r in wb["BOM_chi_tiet"].iter_rows(min_row=9, max_row=2806, max_col=11, values_only=True) if r[0]]

# Kịch bản: (SKU, mét, số lô) — phủ đủ 6 nhóm, bobin + cuộn, nhiều lô
plan = [("PFR-2C-1p5", 3500, 1), ("SFR-2P-1", 2000, 2), ("CTL-1P-A18", 1200, 1), ("SIF-1C-1p5", 10000, 2),
        ("CFR-4C-1", 1500, 1), ("PVC-1C-2p5", 5000, 3), ("PFR-4C-6", 800, 1), ("SFR-4P-2p5", 600, 1)]
plan_m = {s: (m, l) for s, m, l in plan}

dcp, addp, sic, tpc, siw = P["dcPhr"], P["addPhr"], P["siColor"], P["tpColor"], P["siWhite"]
expected = {}
for r in bom:
    sku, code, d_rate, basis = r[0], r[1], float(r[3] or 0), str(r[10] or "")
    if sku not in plan_m:
        continue
    m, lots = plan_m[sku]
    x = dm[sku]
    cores, ins_t, sheath, mpp, pack = x[2], x[4], x[5], x[9], x[10]
    packs = math.ceil(m / mpp)
    pallets = math.ceil(packs / (P["palletW"] if pack == "Bobin gỗ" else P["palletC"]))
    qty = pallets if "pallet" in basis.lower() else (packs if "kiện" in basis.lower() else m / 1000)
    # Phế khởi động theo đúng công thức cột I của BOM_chi_tiet
    AH = lots * cores * (P["startSi"] if ins_t == "SIL" else P["startHD"] if ins_t == "HDPE" else P["startPI"])
    AI = 0 if sheath == "Không" else lots * P["startJ"]
    AK, AL, AM, AN, AO, AP, AQ = x[36], x[37], x[38], x[39], x[40], x[41], x[42]
    colors = {"SI-C0": AK, "SI-C1": AL, "SI-C2": AM, "SI-C3": AN, "SI-C4": AO}
    I = 0.0
    if code == "SIL-9710G": I = AH * AQ
    elif code == "DCBP-50": I = AH * AQ * dcp / 100
    elif code == "SIL-ADD": I = AH * AQ * addp / 100
    elif code in colors: I = AH / (1 + (dcp + addp + sic) / 100) * sic / 100 * colors[code]
    elif code == "SI-W": I = AH * AP / (1 + (dcp + addp + siw) / 100) * siw / 100 if ins_t == "SIL" else 0
    elif code in ("LSZH-S", "PVC-FR", "PVC-S"):
        I = AI / (1 + tpc / 100) * (1 if {"LSZH-S": "LSZH", "PVC-FR": "FR-PVC", "PVC-S": "PVC"}[code] == sheath else 0)
    elif code in ("LS-G", "LS-O", "PF-G", "PF-O", "PV-G", "PV-O", "PV-GR"):
        I = AI / (1 + tpc / 100) * tpc / 100 * (1 if d_rate > 0 else 0)
    elif code == "HDPE-I": I = AH / (1 + tpc / 100) if ins_t == "HDPE" else 0
    elif code == "HD-W": I = AH / (1 + tpc / 100) * tpc / 100 * AP if ins_t == "HDPE" else 0
    elif code == "HD-B": I = AH / (1 + tpc / 100) * tpc / 100 * AM if ins_t == "HDPE" else 0
    elif code == "PVC-I": I = AH / (1 + tpc / 100) if ins_t == "PVC" else 0
    elif code.startswith("PI-C"): I = AH / (1 + tpc / 100) * tpc / 100 * colors["SI-C" + code[-1]] if ins_t == "PVC" else 0
    elif code == "PURGE": I = lots * P["purge"]
    total = d_rate * qty / (1 - mats[code]) + I
    if total:
        expected[code] = expected.get(code, 0) + total

PLAN_NO = f"KH-PARITY-{int(time.time())}"
values = ",".join(f"('{s}',{m},{l})" for s, m, l in plan)
run = psql(f"""
  select set_config('request.jwt.claim.sub', '00000000-0000-0000-0000-0000000000aa', false);
  insert into auth.users(id) values ('00000000-0000-0000-0000-0000000000aa') on conflict do nothing;
  insert into core.profiles(id, full_name, status) values ('00000000-0000-0000-0000-0000000000aa','Kiểm thử MRP','ACTIVE') on conflict do nothing;
  insert into core.user_roles values ('00000000-0000-0000-0000-0000000000aa','SX_KH') on conflict do nothing;
  insert into plan.production_plans (plan_no, period_label, material_need_date, mode) values ('{PLAN_NO}','Kiểm thử', '2026-11-20', 'PRELIMINARY');
  insert into plan.plan_lines (plan_id, product_id, spec_id, qty_m, run_lots, m_per_pack, pack_type, brand)
  select pp.id, p.id, p.current_spec_id, v.m, v.l, p.default_m_per_pack, p.default_pack, p.default_brand
  from (values {values}) v(code, m, l) join master.products p on p.code = v.code, plan.production_plans pp where pp.plan_no = '{PLAN_NO}';
  select plan.run_mrp((select id from plan.production_plans where plan_no = '{PLAN_NO}'));""").splitlines()[-1]
got = json.loads(psql(f"""select json_object_agg(m.code, l.total_issue) from plan.mrp_lines l
  join master.materials m on m.id = l.material_id where l.run_id = '{run}'"""))
worst = 0.0
for code in sorted(set(expected) | set(got)):
    e, g = expected.get(code, 0), float(got.get(code, 0))
    err = abs(g - e) / max(abs(e), 1e-9)
    worst = max(worst, err if abs(g - e) > 1e-4 else 0)
    if abs(g - e) > 1e-4 and err > 1e-6:
        fails += 1
        print(f"SAI MRP {code}: DB {g} ≠ Excel {e}")
print(f"[2] MRP {len(plan)} SKU, {len(expected)} vật tư: sai số tương đối lớn nhất {worst:.2e}")

print("KẾT QUẢ:", "ĐẠT" if fails == 0 else f"{fails} lỗi")
sys.exit(1 if fails else 0)
